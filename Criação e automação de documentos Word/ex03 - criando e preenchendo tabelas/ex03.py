"""
TITULO: CRIAÇÃO E PREENCHIMENTO DE TABELAS
AUTOR: WELLINGTON M SANTOS
DATA: 2026/10/07
OBJETIVO: APRENDER A CRIAR TABELAS, PREENCHER CÉLULAS, ADICIONAR LINHAS DINAMICAMENTE
E APLICAR ESTILOS DE TABELA.
"""

# BIBLIOTECAS
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from loguru import logger

# MAIN
def main():
    logger.info('Configurando a criação do documento word com tabelas.')
    documento = Document()

    # titulo
    logger.info('Inserindo título introdutório')
    documento.add_heading(
        text='Relatório de Vendas - Outubro 2026',
        level=1
    )

    # texto introdutorio
    logger.info('Inserindo texto introdutório')
    documento.add_paragraph(
        text='Segue abaixo o resumo de vendas do mês, gerado automaticamente a partir dos dados do sistema.'
    )

    # dados mockados
    dados_vendas = [
        ('Produto', 'Quantidade', 'Valor Unitário (R$)', 'Total (R$)'),
        ('Notebook Pro 15"', 12, 4500.00, 54000.00),
        ('Mouse Sem Fio', 85, 89.90, 7641.50),
        ('Teclado Mecânico', 34, 350.00, 11900.00),
        ('Monitor 27" 4K', 18, 1890.00, 34020.00),
        ('Headset Gamer', 47, 299.90, 14095.30),
    ]

    # tabela 1x4
    logger.info('Criando tabela')
    tabela = documento.add_table(rows=1, cols=4)
    tabela.style = 'Table Grid'
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER

    # inserir cabeçalhos em negrito
    logger.info('Inserindo cabeçalhos')
    cabecalho = tabela.rows[0].cells
    for i, texto in enumerate(dados_vendas[0]):
        cabecalho[i].text = texto
        for paragrafo in cabecalho[i].paragraphs: # negrito
            for run in paragrafo.runs:
                run.bold = True
                run.font.size = Pt(11)
    
    # popular
    logger.info('Adicionando linhas de dados') 
    for produto, qtd, valor_unit, total in dados_vendas[1:]:
        linha = tabela.add_row().cells
        linha[0].text = produto
        linha[1].text = str(qtd)
        linha[2].text = f'{valor_unit:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
        linha[3].text = f'{total:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

    documento.add_paragraph() # espaço 


    # totais
    logger.info('Inserido resumo de totais')
    total_geral = sum(item[3] for item in dados_vendas[1:])
    paragrafo_total = documento.add_paragraph()
    paragrafo_total.add_run('Total Geral de Vendas: ').bold = True
    paragrafo_total.add_run(f'R$ {total_geral:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'))

    # paragrafo de observação
    documento.add_paragraph(
        text='Observação: Valores gerados automaticamente. Conferir com o sistema de origem'
    ).runs[0].italic = True

    # criando doc
    logger.info('Salvando o documento.')
    documento.save('ex03.docx')
    logger.info('Documento criado com sucesso: ex03.docx')

if __name__ == '__main__':
    main()
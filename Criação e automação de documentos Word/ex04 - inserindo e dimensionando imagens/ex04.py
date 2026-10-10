"""
INSERINDO E DIMENSIONANDO IMAGENS

OBJETIVO:
    APRENDER A ADICIONAR IMAGENS EM DOCUMENTOS WORD COM PYTHON-DOCX,
    CONTROLANDO DIMENSÕES E MANTENDO PROPORÇÃO

AUTOR: WELLINGTON M SANTOS
"""

# BIBLIOTECAS
from docx import Document
from docx.shared import Inches, Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from loguru import logger
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

# LOGO
def criar_logo_exemplo(caminho:str = 'logo_exemplo.png') -> str:
    """
    Cria um logo simples para o exemplo (caso ele não exista)

    Args:
        caminho(str): caminho para o arquivo
    Returns:
        str
    """

    if Path(caminho).exists():
        return caminho

    logger.info(f'Criando logo de exemplo: {caminho}')
    img  = Image.new('RGB', (420,120), color=(30, 64, 175))
    draw = ImageDraw.Draw(img)
    draw.rectangle([10,10,390,110], outline=(255,255,255), width=3)

    try:
        fonte = ImageFont.truetype('DejaVuSans-Bold.ttf', 36)
    except Exception:
        fonte = ImageFont.load_default()

    draw.text(xy=(30,35),
              text="AUTOMACAO DOCS",
               fill=(255,255,255),
                font=fonte)
    img.save(caminho)
    return caminho

# MAIN
def main():
    logger.info('Configurando a criação do documento word com imagens.')
    documento = Document()   

    # titulo
    logger.info('Inserindo título')
    documento.add_heading(text='Relatório com Imagens - Exemplo 04', 
                          level=1)
    
    # texto introdutorio
    logger.info('Inserindo texto de introdução')
    documento.add_paragraph(
        text='Neste exemplo aprendemos a inserir imagens em documentos Word, controlando o tamanho e a posição.'
    )

    # inserindo logo/criando logo
    caminho_logo = criar_logo_exemplo()

    # polegadas
    logger.info('Inserindo logo com largura de 2.5 polegadas')
    paragrafo_logo = documento.add_paragraph()
    paragrafo_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragrafo_logo.add_run()
    run.add_picture(caminho_logo, width=Inches(2.5))
    documento.add_paragraph() # espaço

    # centimetros
    logger.info('Inserindo logo com largura de 8cm')
    documento.add_heading(text='Imagem com tamanho em centímetros', 
                          level=2)
    documento.add_paragraph(
        text='É possível usar Inches() ou Cm(). Ao informar apenas a largura a ' \
        'altura é calculada automaticamente para manter a proporção'
    )
    documento.add_picture(caminho_logo, width=Cm(8))
    documento.add_paragraph() # espaço

    # thumbnail
    logger.info('Inserindo versão thumbnail')
    documento.add_heading(text='Versão reduzida (thumbnail)',
                          level=2)
    p = documento.add_paragraph()
    p.add_run('Logo pequeno: ')
    p.add_run().add_picture(caminho_logo, width=Inches(1.2))

    # observações
    documento.add_heading(text='Observações Importantes',
                          level=2)
    documento.add_paragraph(
        text='Preferir informar apenas width OU height (nunca os dois, a menos que queira distorcer a imagem).',
        style='List Bullet'
    )
    documento.add_paragraph(
        text='Unidades disponíveis: Inches(), Cm(), Pt(), Emu() (unidade interna).',
        style='List Bullet'
    )
    documento.add_paragraph(
        text='O arquivo da imagem deve existir no disco (ou ser um file-like object).',
        style='List Bullet'
    )
    documento.add_paragraph(
        text='Formatos suportados: PNG, JPEG, GIF, BMP, TIFF, etc.',
        style='List Bullet'
    )

    # persistir
    logger.info('Salvando documento')
    documento.save('ex04.docx')
    logger.info('Documento criado com sucesso: ex04.docx')
    logger.info(f'Logo utilizado: {caminho_logo}')


if __name__ == '__main__':
    main()